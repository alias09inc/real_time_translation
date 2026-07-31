"""Extract raw text from presentation-style documents (PDF, PPTX)."""

from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".pptx"}


def extract_document_text(path: Path | str) -> list[str]:
    """Extract text per page/slide from a PDF or PPTX file.

    Args:
        path: Path to a .pdf or .pptx file

    Returns:
        Text per page/slide, in order, with empty pages/slides dropped

    Raises:
        ValueError: If the file extension isn't supported
        RuntimeError: If the required optional dependency isn't installed
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        blocks = _extract_pdf_text(path)
    elif suffix == ".pptx":
        blocks = _extract_pptx_text(path)
    else:
        raise ValueError(
            f"Unsupported document type {suffix!r} for {path}. "
            f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return [block.strip() for block in blocks if block.strip()]


def _extract_pdf_text(path: Path) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "pypdf is not installed. Run: uv sync --extra preload"
        ) from exc

    reader = PdfReader(str(path))
    return [page.extract_text() or "" for page in reader.pages]


def _extract_pptx_text(path: Path) -> list[str]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is not installed. Run: uv sync --extra preload"
        ) from exc

    presentation = Presentation(str(path))
    blocks: list[str] = []
    for slide in presentation.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None and notes_frame.text.strip():
                parts.append(f"[notes] {notes_frame.text}")
        blocks.append("\n".join(parts))
    return blocks
